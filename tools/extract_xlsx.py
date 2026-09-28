# -*- coding: utf-8 -*-
"""
띵타이쿤 가격 계산기 - 엑셀 → web/js/data.js 변환 스크립트

사용법 (프로젝트 루트에서):
    py tools/extract_xlsx.py [엑셀경로]

엑셀 경로를 생략하면 DEFAULT_XLSX 를 사용합니다.
생성 결과는 `window.DEFAULT_DATA = {...};` 형태의 일반 스크립트(web/js/data.js)이며,
file:// 에서도 동작하도록 fetch 없이 <script> 태그로 바로 불러옵니다.
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

# 기본 원본 엑셀 경로 (다른 위치의 엑셀을 쓰려면 인자로 넘기거나 이 값을 바꾸세요)
DEFAULT_XLSX = Path.home() / "Downloads" / "띵타이쿤 요리가격 계산 시트.xlsx"

# 출력 파일 경로 (프로젝트 루트 기준 web/js/data.js)
OUT_PATH = Path(__file__).resolve().parent.parent / "web" / "js" / "data.js"

# 시트 이름 → 카테고리 키 매핑 (시트 이름이 바뀌면 여기를 수정하세요)
SHEETS = {
    "요리": "cooking",
    "공예품": "craft",
}

# 데이터가 시작되는 행 번호 (3행이 헤더, 4행부터 데이터)
FIRST_DATA_ROW = 4

# 열 위치 (1부터 시작: B=2 이름, C=3 최저가, D=4 최고가, E=5 현재가격, I=9 재료)
COL_NAME, COL_MIN, COL_MAX, COL_CUR, COL_ING = 2, 3, 4, 5, 9

# 재료 한 항목 파싱 정규식: "토마토 베이스 2개" → ("토마토 베이스", 2)
ING_RE = re.compile(r"^(.*) (\d+)개$")

# 최신 시세(최저가, 최고가) 덮어쓰기 표 — 엑셀 값보다 우선 적용됩니다.
# 게임 업데이트(밸런스 패치)로 가격 범위가 바뀌면 이 표만 고치고 스크립트를 다시 실행하세요.
# 여기 있는 아이템은 엑셀의 '현재가격'이 옛 범위 기준이라 기본 현재가를 비워 둡니다(게임에서 새로 입력/인식).
# 출처: 띵타이쿤 위키 일반 요리 가격 (2026-09-26 기준)
PRICE_OVERRIDES = {
    # 커먼
    "토마토 스파게티": (171, 570),
    "어니언 링": (203, 677),
    "갈릭 케이크": (149, 499),
    # 노멀
    "삼겹살 토마토 찌개": (403, 1346),
    "삼색 아이스크림": (598, 1995),
    "마늘 양갈비 핫도그": (339, 1131),
    "달콤 시리얼": (510, 1701),
    "로스트 치킨 파이": (422, 1408),
    # 레어
    "스윗 치킨 햄버거": (640, 2134),
    "토마토 파인애플 피자": (609, 2031),
    "양파 수프": (752, 2506),
    "허브 삼겹살 찜": (590, 1968),
    # 에픽
    "토마토 라자냐": (827, 2757),
    "딥 크림 빠네": (760, 2532),
    "트리플 소갈비 꼬치": (852, 2843),
}


def parse_ingredients(text):
    """'A 1개 + B 2개' 형태의 문자열을 [{name, qty}] 리스트로 변환"""
    if not text:
        return []
    result = []
    for part in str(text).split(" + "):
        part = part.strip()
        if not part:
            continue
        m = ING_RE.match(part)
        if not m:
            raise ValueError(f"재료 형식을 해석할 수 없습니다: {part!r}")
        result.append({"name": m.group(1).strip(), "qty": int(m.group(2))})
    return result


def to_num(v):
    """엑셀 숫자(float)를 가능하면 int 로 정리"""
    if v is None:
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx)  # 수식이 아닌 입력값만 쓰므로 data_only 불필요

    items = []
    ingredients = []  # 첫 등장 순서대로 고유 재료명
    for sheet_name, category in SHEETS.items():
        ws = wb[sheet_name]
        for row in range(FIRST_DATA_ROW, ws.max_row + 1):
            name = ws.cell(row, COL_NAME).value
            if name is None or str(name).strip() == "":
                continue
            ing = parse_ingredients(ws.cell(row, COL_ING).value) if category == "cooking" else []
            for it in ing:
                if it["name"] not in ingredients:
                    ingredients.append(it["name"])
            name = str(name).strip()
            lo, hi = to_num(ws.cell(row, COL_MIN).value), to_num(ws.cell(row, COL_MAX).value)
            cur = to_num(ws.cell(row, COL_CUR).value)
            if name in PRICE_OVERRIDES:  # 최신 시세 우선, 옛 현재가는 비움
                lo, hi = PRICE_OVERRIDES[name]
                cur = None
            items.append({
                "category": category,
                "name": name,
                "min": lo,
                "max": hi,
                "current": cur,
                "ingredients": ing,
            })

    # 덮어쓰기 표에 있는데 엑셀에 없는 이름이 있으면 오타일 수 있으므로 경고
    missing = set(PRICE_OVERRIDES) - {it["name"] for it in items}
    if missing:
        print(f"경고: 엑셀에 없는 덮어쓰기 항목: {sorted(missing)}")

    data = {
        "source": xlsx.name,
        "generated": date.today().isoformat(),
        "items": items,
        "ingredients": ingredients,
    }
    body = json.dumps(data, ensure_ascii=False, indent=2)
    header = (
        "// 자동 생성 파일 - 직접 수정하지 마세요.\n"
        "// 재생성: py tools/extract_xlsx.py  (원본 엑셀 → 기본 데이터)\n"
        "// 앱의 '기본값으로 복원' 시 이 데이터로 돌아갑니다.\n"
    )
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUT_PATH.write_text(header + "window.DEFAULT_DATA = " + body + ";\n", encoding="utf-8")
    print(f"OK: {len(items)} items, {len(ingredients)} ingredients -> {OUT_PATH}")


if __name__ == "__main__":
    main()
